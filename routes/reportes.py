"""Reportes PDF y Excel de la cuenta personal."""
from flask import Blueprint, Response, request, session, redirect, url_for
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from io import BytesIO
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from decorators import login_required
from models.persona_model import get_movimientos_filtrados

reportes_bp = Blueprint("reportes", __name__)


def _money(value):
    return "$ {:,.0f}".format(float(value or 0)).replace(",", ".")


def _date(value):
    return str(value or "")[:19]


def _safe_filename(text):
    return str(text or "finex").replace(" ", "_").replace("/", "-")


def _logo_path():
    path = Path(__file__).resolve().parents[1] / "static" / "logo_finex.png"
    return str(path) if path.exists() else None


def _draw_header(p, title, subtitle=""):
    """Dibuja encabezado uniforme para cada hoja del reporte."""
    width, height = letter
    logo = _logo_path()
    if logo:
        try:
            p.drawImage(logo, width - 1.35 * inch, height - 1.05 * inch, width=0.55 * inch, height=0.55 * inch, mask="auto")
        except Exception:
            pass
    p.setFillColor(colors.HexColor("#061c35"))
    p.setFont("Helvetica-Bold", 18)
    p.drawString(0.75 * inch, height - 0.75 * inch, title)
    p.setFont("Helvetica", 9)
    p.setFillColor(colors.HexColor("#536680"))
    if subtitle:
        p.drawString(0.75 * inch, height - 0.98 * inch, subtitle)
    p.setStrokeColor(colors.HexColor("#e5ebf3"))
    p.line(0.75 * inch, height - 1.18 * inch, width - 0.75 * inch, height - 1.18 * inch)
    p.setFillColor(colors.black)
    return height - 1.45 * inch


def _table_page(p, title, movimientos, usuario, correo):
    """Crea una hoja del PDF con una tabla de movimientos."""
    width, height = letter
    y = _draw_header(p, title, f"Usuario: {usuario}  |  Correo: {correo}")
    headers = [("N°", 0.75), ("Fecha", 1.25), ("Categoría", 2.45), ("Descripción", 3.55), ("Monto", 5.25), ("Estado", 6.25)]
    p.setFont("Helvetica-Bold", 8)
    p.setFillColor(colors.HexColor("#071833"))
    for name, x in headers:
        p.drawString(x * inch, y, name)
    y -= 0.12 * inch
    p.setStrokeColor(colors.HexColor("#e5ebf3"))
    p.line(0.75 * inch, y, width - 0.75 * inch, y)
    y -= 0.22 * inch
    p.setFont("Helvetica", 8)
    if not movimientos:
        p.setFillColor(colors.HexColor("#6b7890"))
        p.drawString(0.75 * inch, y, "Sin registros en el rango seleccionado.")
        p.showPage()
        return
    for mov in movimientos:
        if y < 0.75 * inch:
            p.showPage()
            y = _draw_header(p, title, f"Usuario: {usuario}  |  Correo: {correo}")
            p.setFont("Helvetica", 8)
        p.setFillColor(colors.black)
        p.drawString(0.75 * inch, y, str(mov.get("numero_registro") or ""))
        p.drawString(1.25 * inch, y, _date(mov.get("fecha"))[:16])
        p.drawString(2.45 * inch, y, str(mov.get("categoria") or "")[:18])
        p.drawString(3.55 * inch, y, str(mov.get("descripcion") or "")[:28])
        p.drawRightString(5.95 * inch, y, _money(mov.get("valor")))
        p.drawString(6.25 * inch, y, str(mov.get("estado") or ""))
        y -= 0.23 * inch
    p.showPage()


@reportes_bp.route("/reportes/movimientos/pdf")
@login_required
def pdf_movimientos():
    """Exporta un PDF completo: resumen, ingresos y egresos en hojas separadas."""
    desde = request.args.get("desde")
    hasta = request.args.get("hasta")
    usuario_id = session.get("usuario_id")
    movimientos = get_movimientos_filtrados(usuario_id, desde, hasta)
    ingresos = [m for m in movimientos if m.get("tipo") == "ingreso"]
    egresos = [m for m in movimientos if m.get("tipo") == "egreso"]
    total_ingresos = sum(float(m.get("valor") or 0) for m in ingresos if m.get("estado") == "pagada")
    total_egresos = sum(float(m.get("valor") or 0) for m in egresos if m.get("estado") == "pagada")

    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    y = _draw_header(p, "FINEX - Reporte completo personal", f"Usuario: {session.get('usuario_nombre')}  |  Correo: {session.get('usuario_email')}")
    p.setFont("Helvetica-Bold", 12)
    p.drawString(0.75 * inch, y, "Resumen general")
    y -= 0.35 * inch
    p.setFont("Helvetica", 10)
    rango = f"Desde: {desde or 'Inicio'}   Hasta: {hasta or 'Actual'}"
    p.drawString(0.75 * inch, y, rango)
    y -= 0.35 * inch
    p.drawString(0.75 * inch, y, f"Ingresos pagados: {_money(total_ingresos)}")
    y -= 0.25 * inch
    p.drawString(0.75 * inch, y, f"Egresos pagados: {_money(total_egresos)}")
    y -= 0.25 * inch
    p.drawString(0.75 * inch, y, f"Saldo: {_money(total_ingresos - total_egresos)}")
    p.showPage()
    _table_page(p, "Ingresos registrados", ingresos, session.get("usuario_nombre"), session.get("usuario_email"))
    _table_page(p, "Egresos registrados", egresos, session.get("usuario_nombre"), session.get("usuario_email"))
    p.save()
    buffer.seek(0)
    return Response(buffer.read(), mimetype="application/pdf", headers={"Content-Disposition": "attachment; filename=finex_reporte_completo.pdf"})


@reportes_bp.route("/reportes/movimientos/excel")
@login_required
def excel_movimientos():
    """Exporta un Excel completo con hojas separadas para ingresos y egresos."""
    desde = request.args.get("desde")
    hasta = request.args.get("hasta")
    movimientos = get_movimientos_filtrados(session.get("usuario_id"), desde, hasta)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    sheets = [("Ingresos", [m for m in movimientos if m.get("tipo") == "ingreso"]), ("Egresos", [m for m in movimientos if m.get("tipo") == "egreso"])]

    header_fill = PatternFill("solid", fgColor="0B438C")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append(["FINEX - Reporte completo personal"])
    ws.append(["Usuario", session.get("usuario_nombre")])
    ws.append(["Correo", session.get("usuario_email")])
    ws.append(["Desde", desde or "Inicio"])
    ws.append(["Hasta", hasta or "Actual"])
    ws.append([])
    ws.append(["Tipo", "Total pagado", "Registros"])
    for cell in ws[7]:
        cell.fill = header_fill
        cell.font = header_font
    for title, rows in sheets:
        total = sum(float(m.get("valor") or 0) for m in rows if m.get("estado") == "pagada")
        ws.append([title, total, len(rows)])
    for col in range(1, 4):
        ws.column_dimensions[get_column_letter(col)].width = 25

    for title, rows in sheets:
        sheet = wb.create_sheet(title)
        headers = ["N° registro", "Fecha", "Categoría", "Descripción", "Monto", "Estado"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for m in rows:
            sheet.append([m.get("numero_registro"), _date(m.get("fecha")), m.get("categoria"), m.get("descripcion"), float(m.get("valor") or 0), m.get("estado")])
        for col in range(1, len(headers)+1):
            sheet.column_dimensions[get_column_letter(col)].width = 24

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(output.read(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=finex_reporte_completo.xlsx"})
