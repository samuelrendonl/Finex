"""Controlador del panel personal.

Convierte las solicitudes HTTP en llamadas al modelo personal y mantiene
la sesion como filtro principal de seguridad.
"""
from flask import render_template, session, jsonify, request, Response, redirect, url_for, flash
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from openpyxl import Workbook
from models.persona_model import (
    dashboard_data,
    get_movimientos_filtrados,
    insert_movimiento,
    delete_movimiento,
    get_categorias,
    insert_categoria,
    get_movimiento_by_id,
    update_movimiento,
    update_estado_movimiento,
    get_persona_profile,
    update_persona_profile,
)


def _usuario_id():
    return session.get("usuario_id")


def render_dashboard():
    """Renderiza el tablero personal con datos actualizados de registro."""
    perfil = get_persona_profile(_usuario_id())
    nombre_completo = (f"{perfil.get('nombre') or ''} {perfil.get('apellido') or ''}".strip()
                      or session.get("usuario_nombre") or "Usuario")
    correo = perfil.get("email") or session.get("usuario_email") or ""
    return render_template(
        "dashboard.html",
        nombre=nombre_completo,
        correo=correo,
        perfil=perfil,
        dashboard_json=dashboard_data(_usuario_id()),
    )


def obtener_dashboard():
    return jsonify(dashboard_data(_usuario_id()))


def obtener_movimientos():
    desde = request.args.get("desde")
    hasta = request.args.get("hasta")
    tipo = request.args.get("tipo")
    return jsonify(get_movimientos_filtrados(_usuario_id(), desde, hasta, tipo))


def crear_movimiento():
    data = request.get_json() or {}
    movimiento_id = insert_movimiento(data, _usuario_id())
    return jsonify({"success": True, "id": movimiento_id})


def eliminar_movimiento(id):
    delete_movimiento(_usuario_id(), id)
    return jsonify({"success": True})


def obtener_categorias(tipo):
    return jsonify(get_categorias(tipo, _usuario_id()))


def crear_categoria():
    data = request.get_json() or {}
    insert_categoria(data, _usuario_id())
    return jsonify({"success": True})


def obtener_factura(id):
    mov = get_movimiento_by_id(_usuario_id(), id)
    return jsonify(mov or {})


def actualizar_movimiento(id):
    data = request.get_json() or {}
    update_movimiento(_usuario_id(), id, data)
    return jsonify({"success": True})


def actualizar_estado(id, estado):
    update_estado_movimiento(_usuario_id(), id, estado)
    return jsonify({"success": True})


def _fmt(value):
    return "$ {:,.0f}".format(float(value or 0)).replace(",", ".")


def descargar_pdf_movimiento(id):
    """Genera una factura PDF sencilla para un ingreso o egreso personal."""
    mov = get_movimiento_by_id(_usuario_id(), id)
    if not mov:
        return jsonify({"error": "Movimiento no encontrado"}), 404

    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    y = height - 0.8 * inch
    p.setFont("Helvetica-Bold", 18)
    p.drawString(0.75 * inch, y, "FINEX - Factura de movimiento")
    y -= 0.35 * inch
    p.setFont("Helvetica", 10)
    p.drawString(0.75 * inch, y, f"Usuario: {session.get('usuario_nombre')} - {session.get('usuario_email')}")
    y -= 0.55 * inch
    p.setFont("Helvetica-Bold", 12)
    p.drawString(0.75 * inch, y, f"Registro N° {mov['numero_registro']}")
    y -= 0.35 * inch
    p.setFont("Helvetica", 11)
    campos = [
        ("Tipo", mov.get("tipo")),
        ("Categoria", mov.get("categoria")),
        ("Descripcion", mov.get("descripcion")),
        ("Monto", _fmt(mov.get("valor"))),
        ("Fecha", str(mov.get("fecha"))),
        ("Estado", mov.get("estado")),
    ]
    for label, value in campos:
        p.setFont("Helvetica-Bold", 10)
        p.drawString(0.75 * inch, y, f"{label}:")
        p.setFont("Helvetica", 10)
        p.drawString(2.0 * inch, y, str(value or ""))
        y -= 0.28 * inch
    p.showPage()
    p.save()
    buffer.seek(0)
    filename = f"{mov.get('categoria') or 'movimiento'}_{mov['numero_registro']}.pdf".replace(" ", "_")
    return Response(buffer.read(), mimetype="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


def descargar_excel_movimiento(id):
    """Genera un Excel con los datos del movimiento seleccionado."""
    mov = get_movimiento_by_id(_usuario_id(), id)
    if not mov:
        return jsonify({"error": "Movimiento no encontrado"}), 404

    wb = Workbook()
    ws = wb.active
    ws.title = "Movimiento"
    ws.append(["N° registro", "Tipo", "Fecha", "Categoria", "Descripcion", "Monto", "Estado"])
    ws.append([mov["numero_registro"], mov["tipo"], str(mov["fecha"]), mov["categoria"], mov["descripcion"], float(mov["valor"]), mov["estado"]])
    for col in range(1, 8):
        ws.cell(row=1, column=col).font = ws.cell(row=1, column=col).font.copy(bold=True)
        ws.column_dimensions[chr(64 + col)].width = 22

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    fecha = str(mov.get("fecha") or "").split(" ")[0]
    filename = f"{mov.get('categoria') or 'movimiento'}_{mov['numero_registro']}_{fecha}.xlsx".replace(" ", "_")
    return Response(output.read(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={filename}"})


def actualizar_configuracion_personal():
    """Guarda cambios de configuracion de la cuenta personal."""
    try:
        actualizado = update_persona_profile(_usuario_id(), request.form)
        session["usuario_nombre"] = actualizado["nombre"]
        session["usuario_email"] = actualizado["email"]
        flash("Configuración personal actualizada correctamente.", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("persona.dashboard") + "#configuracion")
